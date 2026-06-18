import sys
import os
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill

def main():
    if len(sys.argv) < 3:
        print("Usage: python excel_generator.py <output_path> <selected_mop>", file=sys.stderr)
        sys.exit(1)
        
    output_path = sys.argv[1]
    selected_mop = sys.argv[2]
    
    # Read invoices JSON from stdin
    try:
        invoices_data = sys.stdin.read()
        invoices = json.loads(invoices_data)
    except Exception as e:
        print(f"Error reading/parsing JSON input: {e}", file=sys.stderr)
        sys.exit(1)
        
    # Load template
    template_path = os.path.join(os.path.dirname(__file__), "..", "resources", "report_template.xlsx")
    if not os.path.exists(template_path):
        print(f"Template path not found: {template_path}", file=sys.stderr)
        sys.exit(1)
        
    try:
        wb = load_workbook(template_path)
    except Exception as e:
        print(f"Error loading workbook template: {e}", file=sys.stderr)
        sys.exit(1)
        
    ws1 = wb["Sheet1"]
    ws2 = wb["OW"]
    
    # 1. Update Sheet1 filter cell value (B1)
    # Pivot page filter defaults to (All) if not specified or 'All'
    pivot_filter_val = "(All)"
    if selected_mop and selected_mop.lower() != 'all':
        pivot_filter_val = selected_mop
    ws1["B1"].value = pivot_filter_val
    ws1["B1"].font = Font(name="Calibri", size=11, bold=True)
    
    # 2. Clear old data rows in OW sheet (Row 2 onwards)
    if ws2.max_row > 1:
        # Clear cell values instead of delete_rows to preserve sheet dimensions/styles safely
        for r in range(2, ws2.max_row + 1):
            for c in range(1, 11):
                ws2.cell(r, c).value = None
                ws2.cell(r, c).border = None
                
    # Style definitions for data rows
    thin_border = Border(
        left=Side(style='thin', color='FFD3D3D3'),
        right=Side(style='thin', color='FFD3D3D3'),
        top=Side(style='thin', color='FFD3D3D3'),
        bottom=Side(style='thin', color='FFD3D3D3')
    )
    normal_font = Font(name="Calibri", size=11)
    
    # 3. Add new invoice data rows
    ow_headers = [
        'Case Number',
        'Zip/Postal Code',
        'Product Description',
        'Product Description',
        'Product Item: Product Name: ASP Price',
        'Owner',
        'Customer Name',
        'Remark',
        'Amount',
        'MOP'
    ]
    
    # Ensure Row 1 headers are written correctly
    for col_idx, header in enumerate(ow_headers, 1):
        ws2.cell(1, col_idx).value = header
        
    row_idx = 2
    for inv in invoices:
        prod_desc = inv.get('brand') or 'Atomberg'
        
        # Calculate part description & item price
        items = inv.get('items') or []
        if isinstance(items, str):
            try:
                items = json.loads(items)
            except Exception:
                items = []
                
        part_desc = ", ".join([i.get('description') or i.get('name') or '' for i in items])
        item_price = 0.0
        if len(items) > 0:
            total_rates = sum([float(i.get('rate') or 0) for i in items])
            item_price = total_rates / len(items)
            
        amount_val = float(inv.get('total_amount') or 0)
        
        row_values = [
            inv.get('case_id') or '',
            inv.get('zip_code') or '400001',
            prod_desc,
            part_desc,
            round(item_price, 2),
            inv.get('technician_name') or '',
            inv.get('customer_name') or '',
            inv.get('remark') or '',
            amount_val,
            inv.get('mop') or ''
        ]
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws2.cell(row_idx, col_idx)
            cell.value = val
            cell.font = normal_font
            cell.border = thin_border
            
            # Formatting and Alignment
            if col_idx == 9: # Amount
                cell.number_format = '"₹"#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx == 5: # ASP Price
                cell.number_format = '0.00'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx in [1, 2]: # Numeric codes
                cell.alignment = Alignment(horizontal='left')
                
        row_idx += 1
        
    # 4. Auto column width for OW worksheet
    for col in ws2.columns:
        max_len = 0
        for cell in col:
            val_str = ""
            if cell.value is not None:
                val_str = str(cell.value)
            if cell.number_format and '₹' in cell.number_format:
                val_str = '₹' + val_str
            if len(val_str) > max_len:
                max_len = len(val_str)
        col_letter = col[0].column_letter
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    # 5. Enable filters on OW sheet
    # Determine the end cell range based on written rows
    end_row = max(row_idx - 1, 1)
    ws2.auto_filter.ref = f"A1:J{end_row}"
    
    # 6. Freeze Row 1 on OW sheet
    ws2.freeze_panes = "A2"
    
    # 7. Configure Pivot Tables to refresh on open
    for ws in wb.worksheets:
        for p in ws._pivots:
            p.cache.refreshOnLoad = True
            
    # Save the updated workbook
    try:
        wb.save(output_path)
        print(f"Workbook successfully saved to {output_path}")
    except Exception as e:
        print(f"Error saving workbook: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
