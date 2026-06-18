import os
from openpyxl import load_workbook

files_to_check = [
    "/Users/tanaypatil/Downloads/OW Data.xlsx",
    "/Users/tanaypatil/Downloads/OW Data (1).xlsx",
    "/Users/tanaypatil/Downloads/Electrrolyte March'26 Recivables.xlsx",
    "/Users/tanaypatil/Documents/OW_Report_18-06-2026.xlsx"
]

for filepath in files_to_check:
    if not os.path.exists(filepath):
        print(f"File {filepath} does not exist.")
        continue
        
    print(f"\n==================================================")
    print(f"Inspecting File: {filepath}")
    print(f"==================================================")
    try:
        wb = load_workbook(filepath, read_only=False)
        print("Sheets:", wb.sheetnames)
        for name in wb.sheetnames:
            ws = wb[name]
            pivots = ws._pivots
            print(f"  Sheet '{name}':")
            print(f"    Pivot tables count: {len(pivots)}")
            for idx, p in enumerate(pivots):
                print(f"      Pivot {idx}: name={p.name}, cacheId={p.cacheId}")
                # Print cache source if available
                try:
                    src = p.cache.cacheSource
                    print(f"        CacheSource type: {src.type}")
                    if src.worksheetSource:
                        print(f"        WorksheetSource: sheet={src.worksheetSource.sheet}, ref={src.worksheetSource.ref}")
                except Exception as e:
                    print(f"        Error reading cache: {e}")
    except Exception as err:
        print(f"Error loading {filepath}: {err}")
